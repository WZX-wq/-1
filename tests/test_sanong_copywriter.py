import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from sanong_generation.copywriter import generate_sanong_copywriting
from sanong_generation.errors import TitleGenerationError


def create_title_workbook(root: Path, title: str) -> Path:
    path = root / "titles.xlsx"
    workbook = Workbook()
    workbook.active["B2"] = title
    workbook.save(path)
    return path


class SanongCopywriterTests(unittest.TestCase):
    def test_module_import_does_not_load_pyqt(self):
        script = (
            "import sys; import sanong_generation.copywriter; "
            "assert not any(name.startswith('PyQt6') for name in sys.modules)"
        )
        subprocess.run([sys.executable, "-c", script], check=True)

    def test_generates_copy_label_files_and_updates_database(self):
        with tempfile.TemporaryDirectory() as root:
            title_file = create_title_workbook(Path(root), "闲置土地如何利用")
            with patch(
                "sanong_generation.copywriter.generate_strategy_content",
                return_value=["开场", "介绍闲置土地利用方法"],
            ), patch(
                "sanong_generation.copywriter.generate_topics",
                return_value=["土地利用", "农业政策"],
            ), patch(
                "sanong_generation.copywriter.generate_material_tags",
                side_effect=[["闲置土地", "土地利用"], ["农业政策"]],
            ), patch(
                "sanong_generation.copywriter.segment_title",
                return_value="闲置土地如何利用",
            ), patch(
                "sanong_generation.copywriter.update_video_title_info",
                return_value=True,
            ) as update_info:
                result = generate_sanong_copywriting(title_file, "三农", 321)

            self.assertEqual(result.copy_file.name, "titles_copy.xlsx")
            self.assertEqual(result.label_file.name, "titles_label.xlsx")
            copy_values = [cell.value for cell in load_workbook(result.copy_file).active[2]]
            label_values = [cell.value for cell in load_workbook(result.label_file).active[2]]
            self.assertEqual(
                copy_values,
                ["闲置土地如何利用", "土地利用,农业政策,三农", "开场", "介绍闲置土地利用方法"],
            )
            self.assertEqual(label_values, ["闲置土地,土地利用", "农业政策"])
            update_info.assert_called_once_with(
                321,
                "闲置土地如何利用&&土地利用,农业政策,三农&&开场&&介绍闲置土地利用方法",
                "闲置土地,土地利用&&农业政策",
            )

    def test_empty_copywriting_raises_without_writing_outputs(self):
        with tempfile.TemporaryDirectory() as root:
            root_path = Path(root)
            title_file = create_title_workbook(root_path, "宅基地政策")
            with patch(
                "sanong_generation.copywriter.generate_strategy_content",
                return_value=[],
            ):
                with self.assertRaisesRegex(TitleGenerationError, "文案生成失败"):
                    generate_sanong_copywriting(title_file, "三农", 654)

            self.assertFalse((root_path / "titles_copy.xlsx").exists())
            self.assertFalse((root_path / "titles_label.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
