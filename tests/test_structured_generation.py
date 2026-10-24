import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from ai_tool import doubao_ai_production as ai
from ai_tool.structured_content import normalize_guide_pages, cached_guide_is_valid
from ffmpeg_tool import add_watermark_tool as watermark
from worker.add_watermark_worker import AddWatermarkWorker


ITINERARY = [{'Day1·古城': {'行程安排': '上午古城', '玩点亮点': '城墙',
                             '特色美食': '小吃', '注意事项': '注意安全'}}]
FOOD = [{'第1页·特色美食': {'项目1': {'名称': '【烧饼】', '简介/特色': '酥脆美味',
                                    '价格信息': '10¥/份'}}}]
CUSTOMS = [{'第1页·端午': {'项目1': {'风俗名称': '赛龙舟', '简介寓意': '团结协作',
                                   '具体做法': '划龙舟', '时间地域': '端午'}}}]


class StructuredApiTests(unittest.TestCase):
    def test_cached_broken_copy_is_regenerated_and_valid_copy_is_reused(self):
        self.assertFalse(cached_guide_is_valid('标题&&标签&&"flattened page"', 'food'))
        self.assertTrue(cached_guide_is_valid(
            '标题&&标签&&' + json.dumps(FOOD[0], ensure_ascii=False), 'food'))

    def test_all_guide_apis_preserve_nested_objects_through_excel_roundtrip(self):
        from openpyxl import Workbook, load_workbook
        cases = [(ai.douBao_ai_sights_strategy_content, ITINERARY, 'itinerary'),
                 (ai.douBao_ai_food_strategy_content, FOOD, 'food'),
                 (ai.douBao_ai_festival_customs_content, CUSTOMS, 'customs')]
        with tempfile.TemporaryDirectory() as root:
            for generate, expected, kind in cases:
                for fenced in (False, True):
                    with self.subTest(kind=kind, fenced=fenced):
                        content = json.dumps(expected, ensure_ascii=False)
                        if fenced:
                            content = '```json\n' + content + '\n```'
                        response = SimpleNamespace(choices=[SimpleNamespace(
                            message=SimpleNamespace(content=content))])
                        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(
                            create=lambda **kwargs: response)))
                        with patch.object(ai, '_client', return_value=client):
                            result = generate('测试标题', '测试要求')
                        self.assertEqual(result, expected)
                        workbook = Workbook()
                        workbook.active.append(['测试标题', '标签'] + [
                            json.dumps(page, ensure_ascii=False) for page in result])
                        path = Path(root) / 'copy.xlsx'
                        workbook.save(path)
                        loaded = load_workbook(path)
                        cell = loaded.active.cell(1, 3).value
                        loaded.close()
                        self.assertEqual(normalize_guide_pages(cell, kind), expected)

    def test_malformed_pages_are_rejected_before_rendering(self):
        for kind in ('itinerary', 'food', 'customs'):
            with self.subTest(kind=kind):
                self.assertEqual(normalize_guide_pages(['plain string', {'page': 'bad'}], kind), [])
        malformed = [{'页': {'项目1': {'名称': [], '简介/特色': '测试', '价格信息': ''}}}]
        self.assertEqual(normalize_guide_pages(malformed, 'food'), [])


class WatermarkRegressionTests(unittest.TestCase):
    def test_ffmpeg_error_without_stderr_reaches_worker_error_signal(self):
        with tempfile.TemporaryDirectory() as root:
            image = Path(root) / 'logo.png'
            image.write_bytes(b'image')
            worker = AddWatermarkWorker('input.mp4', 'output.mp4', {
                'watermark_location': ['0', '0'], 'zoom_size': 1, 'pic_alpha': 1,
            }, 2, str(image))
            errors, successes = [], []
            worker.signals.errorOccurred.connect(errors.append)
            worker.signals.finished.connect(successes.append)
            with patch.object(watermark.subprocess, 'run', side_effect=
                              subprocess.CalledProcessError(1, 'ffmpeg', stderr=None)):
                worker.run()
            self.assertEqual(successes, [])
            self.assertEqual(len(errors), 1)
            self.assertIn('FFmpeg图片水印失败', errors[0])
            self.assertNotIn('decode', errors[0])

    def test_missing_watermark_preserves_video_and_emits_clear_message(self):
        with tempfile.TemporaryDirectory() as root:
            source = Path(root) / 'source.mp4'
            output = Path(root) / 'final.mp4'
            source.write_bytes(b'complete generated video')
            worker = AddWatermarkWorker(str(source), str(output), {}, 2, str(Path(root) / 'missing.png'))
            messages = []
            worker.signals.finished.connect(messages.append)
            with patch('worker.add_watermark_worker.add_pic_watermark') as add:
                worker.run()
            self.assertEqual(output.read_bytes(), source.read_bytes())
            self.assertTrue(messages[0].startswith('跳过水印'))
            add.assert_not_called()

    def test_moved_installation_finds_watermark_with_same_filename(self):
        with tempfile.TemporaryDirectory() as root:
            image = Path(root) / 'watermarks' / 'brand.png'
            image.parent.mkdir()
            image.write_bytes(b'image')
            with patch.object(watermark.Path, 'cwd', return_value=Path(root)):
                result = watermark.resolve_watermark_path('Z:/old-computer/brand.png')
            self.assertEqual(result, str(image.resolve()))
