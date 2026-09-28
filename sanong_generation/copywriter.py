from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from ai_tool.ai_text_production import (
    add_suffix_to_file_path,
    segment_title,
    write_arrays_to_excel,
)
from database_tool.video_load_tool import update_video_title_info

from .ai_client import generate_material_tags, generate_strategy_content, generate_topics
from .errors import TitleGenerationError


ProgressCallback = Callable[[str], None]


@dataclass(frozen=True)
class CopywritingFiles:
    copy_file: Path
    label_file: Path


def _notify(progress: ProgressCallback | None, message: str) -> None:
    if progress is not None:
        progress(message)


def _generate_tags(title: str, text: str, title_index: int, part_index: int,
                   progress: ProgressCallback | None, max_retries: int) -> str:
    tags = []
    for attempt in range(1, max_retries + 1):
        _notify(
            progress,
            f"这是第{title_index}个标题，正在ai提取第{part_index}段素材标签 "
            f"(尝试 {attempt}/{max_retries})",
        )
        tags = generate_material_tags(title, text) or []
        if tags:
            break

    if not tags:
        fallback = re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,}", f"{title} {text}")
        tags = list(dict.fromkeys(fallback[:4]))
    cleaned = [str(tag).strip() for tag in tags if str(tag).strip()]
    if not cleaned:
        raise TitleGenerationError(f"第{title_index}个标题第{part_index}段标签生成失败")
    return ",".join(cleaned)


def generate_sanong_copywriting(
    excel_file_path: str | Path,
    account_name: str,
    title_id: int = -1,
    progress: ProgressCallback | None = None,
) -> CopywritingFiles:
    excel_file_path = Path(excel_file_path)
    try:
        workbook = load_workbook(excel_file_path, read_only=True, data_only=True)
        try:
            titles = [
                str(value).strip()
                for (value,) in workbook.active.iter_rows(
                    min_row=2, min_col=2, max_col=2, values_only=True
                )
                if value is not None and str(value).strip()
            ]
        finally:
            workbook.close()
    except Exception as error:
        raise TitleGenerationError(f"读取三农标题文件失败: {error}") from error

    if not titles:
        raise TitleGenerationError("三农标题文件没有可生成的标题")

    max_retries = 3
    generated_rows: list[tuple[list[str], list[str]]] = []
    for title_index, title in enumerate(titles, start=1):
        parts = []
        for attempt in range(1, max_retries + 1):
            _notify(
                progress,
                f"这是第{title_index}个， 正在ai提取该话题文案 "
                f"(尝试 {attempt}/{max_retries})",
            )
            parts = generate_strategy_content(
                title,
                "文案要符合标题，视频文案对观看者有实质性的帮助学习，以通俗易懂的方式讲述,"
                "文案需要控制在大概三百字",
            ) or []
            if parts:
                break
        if not parts:
            raise TitleGenerationError(
                f"第{title_index}个标题文案生成失败，已重试{max_retries}次"
            )

        clean_title = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9\s]", "", title)
        generated_title = segment_title(clean_title)
        topics = [str(item).strip() for item in (generate_topics(",".join(parts)) or [])]
        topics = [item for item in topics if item]
        if account_name:
            topics.append(account_name)
        copy_row = [generated_title, ",".join(topics), *parts]
        label_row = [
            _generate_tags(title, part, title_index, part_index, progress, max_retries)
            for part_index, part in enumerate(parts, start=1)
        ]
        generated_rows.append((copy_row, label_row))

    copy_path = Path(add_suffix_to_file_path(str(excel_file_path), "_copy"))
    label_path = Path(add_suffix_to_file_path(str(excel_file_path), "_label"))
    for path in (copy_path, label_path):
        if path.exists():
            path.unlink()

    for copy_row, label_row in generated_rows:
        write_arrays_to_excel(str(copy_path), copy_row)
        write_arrays_to_excel(str(label_path), label_row)
        if title_id != -1 and not update_video_title_info(
            title_id, "&&".join(copy_row), "&&".join(label_row)
        ):
            raise TitleGenerationError("三农文案标签写入数据库失败")
        _notify(progress, "三农话题文案和标签已经处理完成")

    return CopywritingFiles(copy_file=copy_path, label_file=label_path)
