"""Standalone three-agriculture video composition."""

from __future__ import annotations

import json
import random
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Protocol, Sequence

from openpyxl import Workbook, load_workbook

from database_tool.video_load_tool import (
    cut_video_by_audio_length,
    get_random_video_model,
)
from excel_tool.excel_tool import write_to_old_excel
from ffmpeg_tool.subtitles_audio_tool import get_audio_length

from .errors import MediaGenerationError, TitleGenerationError
from .media_core import (
    add_background_music,
    compose_video_segments,
    generate_audio_subtitles,
)
from .subtitle_style import (
    convert_srt_to_ass,
    create_opening_title,
    remove_first_dialogue_text,
)


EDGE_ROLES = (
    "zh-CN-YunxiNeural",
    "zh-CN-XiaoxiaoNeural",
    "zh-CN-XiaoxuanNeural",
    "zh-CN-XiaoyiNeural",
    "zh-CN-YunjianNeural",
    "zh-CN-YunxiaNeural",
    "zh-CN-YunyangNeural",
    "zh-CN-liaoning-XiaobeiNeural",
    "zh-CN-shaanxi-XiaoniNeural",
)
BAIDU_ROLES = (
    1, 0, 3, 4, 5003, 5118, 106, 5, 4003, 4106, 4115, 5147, 5976,
    5971, 4164, 4176, 4259, 4119, 4105, 4117, 4288, 4192, 4100, 4103,
    4144, 4278, 4143, 4140, 4129, 4149, 4254, 4206, 4147, 4141, 4226,
    4179, 4146, 6567, 4156, 4189, 4194, 4193, 4195, 4196, 4197,
)


@dataclass(frozen=True)
class PreparedSegment:
    audio_path: Path
    ass_path: Path
    audio_length: float


@dataclass(frozen=True)
class PreparedMaterial:
    material_id: int
    video_path: Path


@dataclass(frozen=True)
class SanongVideoResult:
    video_path: Path
    title: str
    tag: str
    description: tuple[str, ...]
    material_ids: tuple[int, ...]


@dataclass(frozen=True)
class SanongComposeOptions:
    screen: int = 0
    random_video_model: bool = False


class MediaBackend(Protocol):
    def prepare_segments(
        self,
        parts: Sequence[str],
        audio_settings: dict,
        settings: dict,
        title: str,
        screen: int,
    ) -> list[PreparedSegment]: ...

    def material_for_segment(
        self,
        tags: Sequence[str],
        used_ids: set[int],
        screen: int,
        audio_length: float,
    ) -> PreparedMaterial: ...

    def render(
        self,
        segments: Sequence[PreparedSegment],
        materials: Sequence[PreparedMaterial],
        intermediate_path: Path,
        pic_time: float,
    ) -> Path: ...

    def add_music(self, intermediate_path: Path, final_path: Path) -> Path: ...


def build_audio_settings(settings: dict) -> dict:
    voice = settings["voice_set"]
    api_type = int(voice["api_type"])
    if api_type == 0:
        speed = str(voice["spd_0"])
        speed = speed if speed.startswith("-") else "+" + speed
        return {
            "api_type": 0,
            "audio_role": EDGE_ROLES[int(voice["role_0"])],
            "speed_value": speed + "%",
            "volume_value": voice["vol_0"],
        }
    return {
        "api_type": 1,
        "audio_role": BAIDU_ROLES[int(voice["role_1"])],
        "speed_value": voice["spd_1"],
        "volume_value": voice["vol_1"],
        "tone_value": voice["pit_1"],
    }


def select_local_sanong_material(
    material_dir: Path,
    used_ids: set[int],
    chooser: Callable[[Sequence[Path]], Path] = random.choice,
) -> PreparedMaterial:
    material_dir = Path(material_dir)
    candidates = [
        path
        for path in sorted(material_dir.glob("*.mp4"), key=lambda item: item.name.lower())
        if path.is_file() and path.stem.isdigit() and int(path.stem) not in used_ids
    ]
    if not candidates:
        available = sum(
            1
            for path in material_dir.glob("*.mp4")
            if path.is_file() and path.stem.isdigit()
        ) if material_dir.is_dir() else 0
        raise TitleGenerationError(
            f"本地三农素材不足：{material_dir} 中有 {available} 个可用 MP4，"
            f"当前视频已使用 {len(used_ids)} 个"
        )
    selected = chooser(candidates)
    return PreparedMaterial(int(selected.stem), selected)


class DefaultMediaBackend:
    def __init__(self, base_dir: Path, music_dir: Path, log: Callable[[str], None]):
        self.base_dir = Path(base_dir)
        self.music_dir = Path(music_dir)
        self.log = log
        self.pic_time = 3.0
        self._title_ass_paths: list[Path] = []

    def prepare_segments(self, parts, audio_settings, settings, title, screen):
        temp_dir = self.base_dir / "tempConcateVideo" / "video_cache_file"
        generated = generate_audio_subtitles(parts, audio_settings, temp_dir, self.log)
        segments = []
        self._title_ass_paths = []
        for index, item in enumerate(generated):
            ass_path = item.srt_path.with_suffix(".ass")
            convert_srt_to_ass(item.srt_path, ass_path, settings["srt_set"])
            if index == 0:
                remove_first_dialogue_text(ass_path)
                title_path = item.srt_path.with_name(item.srt_path.stem + "_title.ass")
                self.pic_time = create_opening_title(
                    settings, screen, title, title_path, item.srt_path
                )
                self._title_ass_paths.append(title_path)
            segments.append(
                PreparedSegment(item.audio_path, ass_path, get_audio_length(str(item.audio_path)))
            )
        return segments

    def material_for_segment(self, tags, used_ids, screen, audio_length):
        del tags, screen
        selected = select_local_sanong_material(
            self.base_dir / "cacheAFRVideo",
            used_ids,
        )
        self.log(f"正在使用本地三农素材：{selected.video_path.name}")
        cut_path = Path(
            cut_video_by_audio_length(
                str(selected.video_path), selected.material_id, audio_length
            ) or ""
        )
        if not cut_path.is_file():
            raise MediaGenerationError(
                f"三农素材 {selected.material_id} 截取后未生成文件"
            )
        return PreparedMaterial(selected.material_id, cut_path)

    def render(self, segments, materials, intermediate_path, pic_time):
        return compose_video_segments(
            [item.video_path for item in materials],
            [item.audio_path for item in segments],
            [item.ass_path for item in segments],
            self._title_ass_paths,
            intermediate_path,
            {"transitionNum": 0, "transitionTime": 0},
            pic_time,
        )

    def add_music(self, intermediate_path, final_path):
        choices = [
            path for path in self.music_dir.iterdir()
            if path.is_file() and path.suffix.lower() == ".mp3"
        ] if self.music_dir.is_dir() else []
        if not choices:
            raise MediaGenerationError(f"背景音乐目录中没有 MP3 文件: {self.music_dir}")
        return add_background_music(
            intermediate_path, random.choice(choices), 100, 40, final_path
        )


def _read_rows(path: Path) -> list[list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            [str(value) for value in row if value is not None]
            for row in workbook.active.iter_rows(values_only=True)
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()


def _append_xlsx(path: Path, values: Sequence[object]) -> Path:
    try:
        if path.exists():
            workbook = load_workbook(path)
        else:
            workbook = Workbook()
        workbook.active.append(list(values))
        workbook.save(path)
        return path
    except PermissionError:
        # WPS/Excel may keep the shared index open. Keep the completed video's
        # metadata in a sidecar workbook instead of turning a successful video
        # into a failed title that will be generated again.
        fallback = path.with_name(
            f"{path.stem}_pending_{datetime.now().strftime('%Y%m%d%H%M%S%f')}{path.suffix}"
        )
        pending = Workbook()
        pending.active.append(list(values))
        pending.save(fallback)
        return fallback


def _safe_video_name(title: str) -> str:
    return re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", title) or "三农视频"


class SanongComposer:
    def __init__(
        self,
        media: MediaBackend | None = None,
        options: SanongComposeOptions | None = None,
        log: Callable[[str], None] | None = None,
        base_dir: Path | None = None,
        music_dir: Path | None = None,
    ):
        self.options = options or SanongComposeOptions()
        self.log = log or (lambda _message: None)
        self.base_dir = Path(base_dir or Path.cwd())
        self.media = media or DefaultMediaBackend(
            self.base_dir,
            Path(music_dir or self.base_dir / "cacheMusic1"),
            self.log,
        )

    def _effective_settings(self, settings: dict) -> dict:
        if not self.options.random_video_model:
            return settings
        model = get_random_video_model()
        if not model or not model.get("model_set"):
            raise TitleGenerationError("没有可用的随机视频模板")
        random_settings = json.loads(model["model_set"])
        merged = dict(settings)
        merged.update(random_settings)
        return merged

    def _cleanup_title_artifacts(
        self,
        segments: Sequence[PreparedSegment],
        materials: Sequence[PreparedMaterial],
    ) -> None:
        roots = tuple(
            root.resolve()
            for root in (
                self.base_dir / "tempConcateVideo",
                self.base_dir / "cacheCutVideo",
            )
        )
        paths = [
            path
            for segment in segments
            for path in (
                segment.audio_path,
                segment.ass_path,
                segment.ass_path.with_suffix(".srt"),
            )
        ]
        paths.extend(material.video_path for material in materials)
        paths.extend(getattr(self.media, "_title_ass_paths", ()))

        for candidate in paths:
            path = Path(candidate)
            try:
                resolved = path.resolve()
                if not resolved.is_file():
                    continue
                if not any(resolved == root or root in resolved.parents for root in roots):
                    continue
                resolved.unlink()
                for root in roots:
                    if resolved == root or root not in resolved.parents:
                        continue
                    parent = resolved.parent
                    while parent != root:
                        try:
                            parent.rmdir()
                        except OSError:
                            break
                        parent = parent.parent
                    break
            except OSError:
                self.log(f"临时文件清理失败，稍后可手动删除：{path}")

        title_ass_paths = getattr(self.media, "_title_ass_paths", None)
        if isinstance(title_ass_paths, list):
            title_ass_paths.clear()

    def compose(
        self,
        copy_file: Path,
        label_file: Path | None,
        output_dir: Path,
        account_name: str,
        settings: dict,
    ) -> list[SanongVideoResult]:
        del account_name  # Kept in the public API for GUI/service symmetry.
        settings = self._effective_settings(settings)
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        copy_rows = _read_rows(Path(copy_file))
        label_rows = _read_rows(Path(label_file)) if label_file and Path(label_file).is_file() else []
        audio_settings = build_audio_settings(settings)
        results = []

        for row_index, row in enumerate(copy_rows):
            if len(row) < 5:
                raise TitleGenerationError("三农文案少于三个片段")
            title, tag, *parts = (item.strip() for item in row)
            labels = label_rows[row_index] if row_index < len(label_rows) else []
            self.log(f"开始合成三农视频：{title}")
            segments: list[PreparedSegment] = []
            materials: list[PreparedMaterial] = []
            intermediate_path: Path | None = None
            try:
                segments = self.media.prepare_segments(
                    parts, audio_settings, settings, title, self.options.screen
                )
                if len(segments) != len(parts):
                    raise MediaGenerationError("生成的音频字幕片段数量不完整")

                used_ids: set[int] = set()
                for index, segment in enumerate(segments):
                    segment_tags = [
                        item.strip()
                        for item in (
                            labels[index] if index < len(labels) else ""
                        ).split(",")
                        if item.strip()
                    ]
                    material = self.media.material_for_segment(
                        segment_tags,
                        used_ids,
                        self.options.screen,
                        segment.audio_length,
                    )
                    if material.material_id in used_ids:
                        raise MediaGenerationError(
                            f"素材 {material.material_id} 被重复使用"
                        )
                    used_ids.add(material.material_id)
                    materials.append(material)

                safe_name = _safe_video_name(title)
                intermediate_path = output_dir / f"{safe_name}_0.mp4"
                final_path = output_dir / f"{safe_name}.mp4"
                rendered = self.media.render(
                    segments,
                    materials,
                    intermediate_path,
                    float(getattr(self.media, "pic_time", 3.0)),
                )
                if Path(rendered) != intermediate_path or not intermediate_path.is_file():
                    raise MediaGenerationError(
                        f"视频合成未生成预期文件: {intermediate_path}"
                    )
                produced = Path(self.media.add_music(intermediate_path, final_path))
                if produced != final_path or not final_path.is_file():
                    raise MediaGenerationError(
                        f"添加音乐未生成预期文件: {final_path}"
                    )

                index_path = output_dir / "video_id.xlsx"
                try:
                    saved_index = _append_xlsx(
                        index_path,
                        [title, *(str(item.video_path) for item in materials)],
                    )
                    if saved_index != index_path:
                        self.log(
                            f"video_id.xlsx 正被其他程序占用，已暂存为：{saved_index.name}"
                        )
                except PermissionError as error:
                    self.log(f"video_id.xlsx 写入失败，已保留视频：{error}")

                try:
                    write_to_old_excel(
                        str(output_dir / "1.xls"),
                        {
                            "标题": title.replace("|", ""),
                            "标签": tag,
                            "文件路径": str(final_path).replace("/", "\\"),
                            "描述": ",".join(parts),
                            "封面路径": "",
                            "图片集合": "",
                            "关联商品": "",
                            "分区分类": "旅游",
                            "备注": "",
                        },
                    )
                except PermissionError as error:
                    self.log(f"1.xls 写入失败，已保留视频：{error}")

                results.append(
                    SanongVideoResult(
                        final_path,
                        title.replace("|", ""),
                        tag,
                        tuple(parts),
                        tuple(item.material_id for item in materials),
                    )
                )
                self.log(f"生成完成一个三农视频：{final_path}")
            finally:
                if intermediate_path is not None:
                    intermediate_path.unlink(missing_ok=True)
                self._cleanup_title_artifacts(segments, materials)
        return results
