from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


HEADER_NAMES = {"标题", "title"}


def read_title_workbook(path: str | Path) -> list[str]:
    """Read unique, non-empty titles from the first column of an xlsx file."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        titles: list[str] = []
        seen: set[str] = set()

        for row_number, row in enumerate(
            sheet.iter_rows(min_col=1, max_col=1, values_only=True), start=1
        ):
            value = row[0]
            if value is None:
                continue

            title = str(value).strip()
            if not title:
                continue
            if row_number == 1 and title.casefold() in HEADER_NAMES:
                continue
            if title in seen:
                continue

            seen.add(title)
            titles.append(title)

        return titles
    finally:
        workbook.close()


__all__ = ["read_title_workbook"]
